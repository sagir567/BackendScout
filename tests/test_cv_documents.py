from pathlib import Path

import pytest
from docx import Document

from backend_scout import cv_documents
from backend_scout.cv_documents import (
    MIN_CV_PAGE_FILL_RATIO,
    PdfLayoutMetrics,
    _select_valid_layout,
    create_cv_docx,
    render_checked_cv_artifacts,
    resolve_libreoffice_executable,
)
from backend_scout.cv_tailoring import build_evidence_only_draft
from backend_scout.models import (
    CareerEvidence,
    CvStyle,
    TailoredBullet,
    TailoredCv,
    TailoredProject,
)


def make_evidence() -> CareerEvidence:
    return CareerEvidence.model_validate(
        {
            "identity": {
                "full_name": "Test Candidate",
                "email": "candidate@example.com",
                "links": [
                    "https://github.com/test-candidate",
                    "https://github.com/test-candidate/cpp",
                    "https://github.com/example-organization",
                    "https://www.linkedin.com/in/test-candidate",
                ],
            },
            "summary": "Backend developer with Python experience.",
            "skills": [{"category": "Backend", "items": ["Python", "FastAPI"]}],
            "experience": [
                {
                    "id": "exp_1",
                    "organization": "Example Company",
                    "title": "Backend Developer",
                    "start_date": "2024-01",
                    "end_date": "Present",
                    "bullets": [
                        {
                            "id": "exp_1_bullet_1",
                            "text": "Built Python APIs for internal tools.",
                        }
                    ],
                }
            ],
        }
    )


def test_create_cv_docx_writes_a_readable_word_document(tmp_path: Path) -> None:
    evidence = make_evidence()
    output_path = tmp_path / "cv_draft.docx"

    create_cv_docx(evidence, build_evidence_only_draft(evidence), output_path)

    assert output_path.is_file()
    assert output_path.stat().st_size > 0


def test_create_cv_docx_uses_labeled_hyperlinks_and_hides_headline(tmp_path: Path) -> None:
    evidence = make_evidence()
    output_path = tmp_path / "cv_draft.docx"
    draft = TailoredCv(
        headline="Backend Engineer | Python",
        summary="Backend developer with Python experience.",
        summary_evidence_ids=["exp_1_bullet_1"],
    )

    create_cv_docx(evidence, draft, output_path, CvStyle())

    document = Document(output_path)
    visible_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    hyperlink_targets = [relationship.target_ref for relationship in document.part.rels.values()]
    assert "https://github.com/test-candidate" not in visible_text
    assert "Backend Engineer | Python" not in visible_text
    assert "https://github.com/test-candidate" in hyperlink_targets
    assert "https://www.linkedin.com/in/test-candidate" in hyperlink_targets
    assert "https://github.com/test-candidate/cpp" not in hyperlink_targets
    assert "https://github.com/example-organization" not in hyperlink_targets


def test_create_cv_docx_uses_per_link_label(tmp_path: Path) -> None:
    evidence = make_evidence()
    output_path = tmp_path / "cv_draft.docx"

    create_cv_docx(
        evidence,
        build_evidence_only_draft(evidence),
        output_path,
        CvStyle(link_labels={"https://github.com/test-candidate": "Personal GitHub"}),
    )

    document = Document(output_path)
    xml = document.part.element.xml
    assert "Personal GitHub" in xml
    assert "https://github.com/test-candidate" not in "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )


def test_project_repository_link_is_rendered_only_with_selected_project(tmp_path: Path) -> None:
    evidence_data = make_evidence().model_dump(mode="json")
    evidence_data["projects"] = [
        {
            "id": "project_cpp",
            "name": "C++ Practice",
            "link": "https://github.com/test-candidate/cpp",
            "selection_keywords": ["C++"],
            "bullets": [
                {
                    "id": "project_cpp_bullet",
                    "text": "Built coursework projects in C++.",
                }
            ],
        }
    ]
    evidence = CareerEvidence.model_validate(evidence_data)
    draft = TailoredCv(
        projects=[
            TailoredProject(
                evidence_id="project_cpp",
                bullets=[
                    TailoredBullet(
                        text="Built coursework projects in C++.",
                        evidence_ids=["project_cpp_bullet"],
                    )
                ],
            )
        ]
    )
    output_path = tmp_path / "cpp_cv.docx"

    create_cv_docx(evidence, draft, output_path, CvStyle())

    document = Document(output_path)
    hyperlink_targets = [relationship.target_ref for relationship in document.part.rels.values()]
    assert "https://github.com/test-candidate/cpp" in hyperlink_targets


def test_libreoffice_resolver_supports_launchd_without_shell_path(tmp_path: Path) -> None:
    executable = tmp_path / "soffice"
    executable.write_text("fixture", encoding="utf-8")
    executable.chmod(0o755)

    resolved = resolve_libreoffice_executable(
        configured_path="",
        lookup=lambda command: None,
        fallback_paths=[executable],
    )

    assert resolved == str(executable)


def test_layout_selection_requires_one_page_with_at_least_ninety_percent_fill(tmp_path: Path) -> None:
    short_docx = tmp_path / "short.docx"
    short_pdf = tmp_path / "short.pdf"
    valid_docx = tmp_path / "valid.docx"
    valid_pdf = tmp_path / "valid.pdf"
    selected, selected_docx, selected_pdf = _select_valid_layout(
        [
            (PdfLayoutMetrics(page_count=1, content_fill_ratio=0.89), short_docx, short_pdf),
            (PdfLayoutMetrics(page_count=1, content_fill_ratio=0.93), valid_docx, valid_pdf),
            (PdfLayoutMetrics(page_count=2, content_fill_ratio=0.0), tmp_path / "two.docx", tmp_path / "two.pdf"),
        ]
    )

    assert selected.content_fill_ratio >= MIN_CV_PAGE_FILL_RATIO
    assert selected_docx == valid_docx
    assert selected_pdf == valid_pdf


def test_layout_selection_rejects_sparse_or_multi_page_candidates(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="at least 90%"):
        _select_valid_layout(
            [
                (PdfLayoutMetrics(page_count=1, content_fill_ratio=0.89), tmp_path / "short.docx", tmp_path / "short.pdf"),
                (PdfLayoutMetrics(page_count=2, content_fill_ratio=0.0), tmp_path / "two.docx", tmp_path / "two.pdf"),
            ]
        )


def test_render_uses_vertical_spacing_fallback_without_changing_typography(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    observed_scales: list[tuple[float, float | None]] = []

    def fake_create(
        evidence: CareerEvidence,
        draft: TailoredCv,
        output_path: Path,
        style: CvStyle,
        layout_scale: float = 1.0,
        vertical_scale: float | None = None,
    ) -> Path:
        observed_scales.append((layout_scale, vertical_scale))
        output_path.write_text("docx", encoding="utf-8")
        return output_path

    def fake_convert(docx_path: Path, pdf_path: Path) -> Path:
        pdf_path.write_text("pdf", encoding="utf-8")
        return pdf_path

    def fake_measure(pdf_path: Path) -> PdfLayoutMetrics:
        if pdf_path.name == "spacing-candidate-0.pdf":
            return PdfLayoutMetrics(page_count=1, content_fill_ratio=0.91)
        if pdf_path.name.startswith("spacing-candidate-"):
            return PdfLayoutMetrics(page_count=2, content_fill_ratio=0.0)
        return PdfLayoutMetrics(page_count=1, content_fill_ratio=0.88)

    monkeypatch.setattr(cv_documents, "create_cv_docx", fake_create)
    monkeypatch.setattr(cv_documents, "convert_docx_to_pdf", fake_convert)
    monkeypatch.setattr(cv_documents, "measure_pdf_layout", fake_measure)

    output_docx = tmp_path / "result.docx"
    output_pdf = tmp_path / "result.pdf"
    metrics = render_checked_cv_artifacts(
        make_evidence(),
        build_evidence_only_draft(make_evidence()),
        output_docx,
        output_pdf,
    )

    assert metrics.content_fill_ratio == 0.91
    assert output_docx.read_text(encoding="utf-8") == "docx"
    assert output_pdf.read_text(encoding="utf-8") == "pdf"
    fallback_scales = [scales for scales in observed_scales if scales[1] is not None]
    assert fallback_scales
    assert fallback_scales[0][0] == 1.0
    assert fallback_scales[0][1] == 1.02
