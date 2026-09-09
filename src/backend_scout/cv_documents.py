import os
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

from backend_scout.models import CareerEvidence, CvStyle, TailoredCv

MIN_CV_PAGE_FILL_RATIO = 0.90
_LAYOUT_SCALES = (1.0, 1.03, 1.06, 1.09, 1.12, 1.15, 1.18, 1.21, 1.24)
_VERTICAL_MARGIN_INCHES = 0.55


@dataclass(frozen=True)
class PdfLayoutMetrics:
    page_count: int
    content_fill_ratio: float


def render_checked_cv_artifacts(
    evidence: CareerEvidence,
    draft: TailoredCv,
    docx_path: Path,
    pdf_path: Path,
    style: CvStyle | None = None,
) -> PdfLayoutMetrics:
    """Render exactly one well-filled CV page or fail without writing final artifacts."""
    style = style or CvStyle()
    candidates: list[tuple[PdfLayoutMetrics, Path, Path]] = []
    with tempfile.TemporaryDirectory(prefix="backendscout-cv-layout-") as temporary_directory:
        temporary_root = Path(temporary_directory)
        for index, scale in enumerate(_LAYOUT_SCALES):
            candidate_docx = temporary_root / f"candidate-{index}.docx"
            candidate_pdf = temporary_root / f"candidate-{index}.pdf"
            create_cv_docx(evidence, draft, candidate_docx, style, layout_scale=scale)
            convert_docx_to_pdf(candidate_docx, candidate_pdf)
            metrics = measure_pdf_layout(candidate_pdf)
            if metrics.page_count == 1:
                candidates.append((metrics, candidate_docx, candidate_pdf))

        metrics, selected_docx, selected_pdf = _select_valid_layout(candidates)
        docx_path.parent.mkdir(parents=True, exist_ok=True)
        pdf_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(selected_docx, docx_path)
        shutil.copy2(selected_pdf, pdf_path)
        return metrics


def measure_pdf_layout(pdf_path: Path) -> PdfLayoutMetrics:
    """Measure how far visible text reaches through a one-page PDF's usable height."""
    import pdfplumber

    with pdfplumber.open(pdf_path) as document:
        if len(document.pages) != 1:
            return PdfLayoutMetrics(page_count=len(document.pages), content_fill_ratio=0.0)
        page = document.pages[0]
        words = page.extract_words()
        if not words:
            raise ValueError("Rendered CV PDF contains no readable text")
        top = _VERTICAL_MARGIN_INCHES * 72
        bottom = page.height - top
        content_bottom = max(float(word["bottom"]) for word in words)
        fill_ratio = (content_bottom - top) / (bottom - top)
        return PdfLayoutMetrics(page_count=1, content_fill_ratio=max(0.0, min(1.0, fill_ratio)))


def _select_valid_layout(
    candidates: list[tuple[PdfLayoutMetrics, Path, Path]],
) -> tuple[PdfLayoutMetrics, Path, Path]:
    valid = [
        candidate
        for candidate in candidates
        if candidate[0].page_count == 1 and candidate[0].content_fill_ratio >= MIN_CV_PAGE_FILL_RATIO
    ]
    if not valid:
        highest_fill = max((metrics.content_fill_ratio for metrics, _, _ in candidates), default=0.0)
        raise ValueError(
            "Could not render a one-page CV that fills at least "
            f"{MIN_CV_PAGE_FILL_RATIO:.0%} of the usable page height "
            f"(best one-page fill: {highest_fill:.0%})."
        )
    return max(valid, key=lambda candidate: candidate[0].content_fill_ratio)


def create_cv_docx(
    evidence: CareerEvidence,
    draft: TailoredCv,
    output_path: Path,
    style: CvStyle | None = None,
    layout_scale: float = 1.0,
) -> Path:
    """Render a single-column, evidence-audited CV using a compact technical layout."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(_VERTICAL_MARGIN_INCHES)
    section.bottom_margin = Inches(_VERTICAL_MARGIN_INCHES)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10 * layout_scale)
    normal.paragraph_format.space_after = Pt(2 * layout_scale)
    normal.paragraph_format.line_spacing = layout_scale

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name.add_run(evidence.identity.full_name)
    name_run.bold = True
    name_run.font.name = "Calibri"
    name_run.font.size = Pt(17 * layout_scale)
    name_run.font.color.rgb = RGBColor(11, 37, 69)
    name.paragraph_format.space_after = Pt(2)

    style = style or CvStyle()
    contact = document.add_paragraph()
    contact_values = [evidence.identity.email, evidence.identity.phone, evidence.identity.location]
    _add_contact_line(contact, contact_values, evidence.identity.links, style.link_labels)
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(6 * layout_scale)

    if style.include_headline and draft.headline:
        headline = document.add_paragraph(draft.headline)
        headline.alignment = WD_ALIGN_PARAGRAPH.CENTER
        headline.runs[0].italic = True
        headline.paragraph_format.space_after = Pt(8)

    if draft.summary:
        _add_heading(document, "SUMMARY", layout_scale)
        document.add_paragraph(draft.summary)

    if draft.skills:
        _add_heading(document, "TECHNICAL SKILLS", layout_scale)
        for group in draft.skills:
            skill_line = document.add_paragraph()
            category = skill_line.add_run(f"{group.category}: ")
            category.bold = True
            skill_line.add_run(", ".join(group.items))

    experience_by_id = {item.id: item for item in evidence.experience}
    if draft.experience:
        _add_heading(document, "EXPERIENCE", layout_scale)
        for tailored in draft.experience:
            source = experience_by_id[tailored.evidence_id]
            _add_role_heading(
                document, source.title, source.organization, source.start_date, source.end_date, layout_scale
            )
            if source.location:
                location = document.add_paragraph(source.location)
                location.runs[0].italic = True
                location.paragraph_format.space_after = Pt(1)
            _add_bullets(document, [bullet.text for bullet in tailored.bullets], layout_scale)

    project_by_id = {item.id: item for item in evidence.projects}
    if draft.projects:
        _add_heading(document, "PROJECTS", layout_scale)
        for tailored in draft.projects:
            source = project_by_id[tailored.evidence_id]
            project = document.add_paragraph()
            project.add_run(source.name).bold = True
            if source.link:
                project.add_run(" | ")
                _add_hyperlink(project, _link_label(source.link, style.link_labels), source.link)
            _add_bullets(document, [bullet.text for bullet in tailored.bullets], layout_scale)

    if draft.education:
        _add_heading(document, "EDUCATION", layout_scale)
        for item in draft.education:
            line = document.add_paragraph()
            line.add_run(item.credential).bold = True
            line.add_run(f", {item.institution}")
            if item.end_date:
                line.add_run(f" | {item.end_date}")

    if draft.publications:
        _add_heading(document, "PUBLICATIONS", layout_scale)
        for item in draft.publications:
            line = document.add_paragraph()
            line.add_run(item.title).bold = True
            line.add_run(f", {item.venue} ({item.year})")

    document.save(output_path)
    return output_path


def convert_docx_to_pdf(docx_path: Path, pdf_path: Path) -> Path:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice is required to generate a PDF CV draft")

    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="backendscout-lo-") as profile:
        environment = os.environ | {"HOME": profile, "TMPDIR": profile}
        result = subprocess.run(
            [
                soffice,
                "--headless",
                f"-env:UserInstallation=file://{profile}",
                "--convert-to",
                "pdf",
                "--outdir",
                str(pdf_path.parent),
                str(docx_path),
            ],
            check=False,
            capture_output=True,
            text=True,
            env=environment,
        )
    generated_pdf = pdf_path.parent / f"{docx_path.stem}.pdf"
    if result.returncode != 0 or not generated_pdf.is_file():
        raise RuntimeError(f"LibreOffice PDF conversion failed: {result.stderr.strip()}")
    if generated_pdf != pdf_path:
        generated_pdf.replace(pdf_path)
    return pdf_path


def _add_heading(document, text: str, layout_scale: float = 1.0) -> None:
    from docx.shared import Pt, RGBColor

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6 * layout_scale)
    paragraph.paragraph_format.space_after = Pt(2 * layout_scale)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(10.5 * layout_scale)
    run.font.color.rgb = RGBColor(31, 77, 120)


def _add_role_heading(
    document,
    title: str,
    organization: str,
    start_date: str,
    end_date: str,
    layout_scale: float = 1.0,
) -> None:
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.shared import Inches, Pt

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3 * layout_scale)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.4), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES
    )
    paragraph.add_run(f"{title} | {organization}").bold = True
    paragraph.add_run(f"\t{start_date} - {end_date}")


def _add_bullets(document, bullets: list[str], layout_scale: float = 1.0) -> None:
    from docx.shared import Pt

    for bullet in bullets:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(bullet)
        paragraph.paragraph_format.space_after = Pt(0 * layout_scale)


def _add_contact_line(
    paragraph,
    values: list[str | None],
    links: list[str],
    custom_labels: dict[str, str],
) -> None:
    visible_values = [value for value in values if value]
    for index, value in enumerate(visible_values):
        if index:
            paragraph.add_run(" | ")
        paragraph.add_run(value)
    for link in links:
        if visible_values or link != links[0]:
            paragraph.add_run(" | ")
        _add_hyperlink(paragraph, _link_label(link, custom_labels), link)


def _add_hyperlink(paragraph, label: str, url: str) -> None:
    """Add a visible label backed by a clickable external Word hyperlink."""
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    relationship_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([color, underline])
    text = OxmlElement("w:t")
    text.text = label
    run.extend([run_properties, text])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _link_label(url: str, custom_labels: dict[str, str]) -> str:
    if label := custom_labels.get(url):
        return label
    domain = urlparse(url).netloc.casefold().removeprefix("www.")
    if domain == "github.com":
        return "GitHub"
    if domain == "linkedin.com":
        return "LinkedIn"
    if domain in {"itch.io", "sagir567.itch.io"} or domain.endswith(".itch.io"):
        return "Project page"
    return domain or "Website"
